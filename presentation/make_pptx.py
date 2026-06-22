#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
タスク管理アプリの発表用スライド（10枚）を生成し、4枚目にデモ動画を埋め込む。
デザイン: 明るくシンプル（白基調・スレート/ブルー）。日本語フォントは Meiryo を指定。
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
FONT = "Meiryo"

INK = RGBColor(0x1E, 0x29, 0x3B)
MUTED = RGBColor(0x64, 0x74, 0x8B)
ACCENT = RGBColor(0x25, 0x63, 0xEB)
LIGHT = RGBColor(0xF1, 0xF5, 0xF9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BORDER = RGBColor(0xE2, 0xE8, 0xF0)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = prs.slide_width, prs.slide_height


def _set_ea(run, name):
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", name)


def run_fmt(run, size, bold=False, color=INK, italic=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = FONT
    _set_ea(run, FONT)


def textbox(slide, left, top, width, height, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    return tf


def rect(slide, left, top, width, height, fill, line=None, line_w=None):
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    sp.fill.solid()
    sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = line_w or Pt(1)
    sp.shadow.inherit = False
    return sp


def add_para(tf, text, size, bold=False, color=INK, bullet=False, space_after=10, first=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_after = Pt(space_after)
    p.line_spacing = 1.15
    if bullet:
        r = p.add_run(); r.text = "●  "; run_fmt(r, size * 0.7, bold=True, color=ACCENT)
    r = p.add_run(); r.text = text; run_fmt(r, size, bold=bold, color=color)
    return p


def content_header(slide, no, title):
    # 上部アクセントバー
    rect(slide, 0, 0, SW, Inches(0.18), ACCENT)
    tf = textbox(slide, Inches(0.7), Inches(0.45), Inches(12), Inches(1.0))
    add_para(tf, title, 30, bold=True, color=INK, first=True)
    # タイトル下の細いルール
    rect(slide, Inches(0.72), Inches(1.45), Inches(2.2), Pt(3), ACCENT)
    # フッタ
    ftf = textbox(slide, Inches(0.7), SH - Inches(0.5), Inches(10), Inches(0.4))
    add_para(ftf, "タスク管理アプリ", 11, color=MUTED, first=True)
    ntf = textbox(slide, SW - Inches(1.2), SH - Inches(0.5), Inches(0.8), Inches(0.4))
    pno = add_para(ntf, f"{no} / 10", 11, color=MUTED, first=True)
    pno.alignment = PP_ALIGN.RIGHT


def body_box(slide):
    return textbox(slide, Inches(0.9), Inches(1.8), Inches(11.5), Inches(5.0))


# ============ Slide 1: タイトル ============
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, SW, SH, WHITE)
rect(s, 0, 0, Inches(0.35), SH, ACCENT)          # 左アクセント帯
rect(s, 0, SH - Inches(0.35), SW, Inches(0.35), INK)
tf = textbox(s, Inches(1.0), Inches(2.2), Inches(11.3), Inches(2.2))
add_para(tf, "【作品タイトル】", 46, bold=True, color=INK, first=True, space_after=6)
add_para(tf, "Next.js（App Router）製 マルチテナント対応 タスク管理アプリ", 22, color=ACCENT)
info = textbox(s, Inches(1.0), Inches(4.9), Inches(11), Inches(1.4))
add_para(info, "学生番号：【学生番号】", 22, color=INK, first=True, space_after=8)
add_para(info, "氏名　　：【氏名】", 22, color=INK)

# ============ Slide 2: テーマ選定理由 ============
s = prs.slides.add_slide(BLANK)
content_header(s, 2, "なぜこのテーマを選んだのか")
tf = body_box(s)
add_para(tf, "日々のタスク管理を「自分の手で」効率化できるアプリを作りたかった。", 22, bold=True, first=True, space_after=16)
add_para(tf, "紙やメモアプリでは、優先度・期限・進捗が一目で分からず管理しきれなかった。", 21, bullet=True)
add_para(tf, "Web アプリをフロントからバックエンドまで一気通貫で作れたら面白いと思った。", 21, bullet=True)
add_para(tf, "「動くだけ」でなく、実際のサービスのように使える品質に挑戦してみたかった。", 21, bullet=True)
add_para(tf, "ドラッグ＆ドロップなど、触っていて気持ちいい操作を自分で実装したかった。", 21, bullet=True)

# ============ Slide 3: 作品概要 ============
s = prs.slides.add_slide(BLANK)
content_header(s, 3, "作品概要 ― できること")
tf = body_box(s)
add_para(tf, "ユーザー登録／ログインして、自分専用のタスクを管理できる Web アプリ。", 22, bold=True, first=True, space_after=16)
add_para(tf, "タスクの作成・編集・削除（CRUD）と、未着手／進行中／完了の状態管理", 21, bullet=True)
add_para(tf, "ドラッグ＆ドロップのカンバンボードで、カードを動かして状態変更", 21, bullet=True)
add_para(tf, "タグ付け・優先度・期限の設定と、期限超過の自動ハイライト", 21, bullet=True)
add_para(tf, "タイトル／説明のインクリメンタル検索、ステータス別タブと件数サマリー", 21, bullet=True)
add_para(tf, "優先度 → 期限 → 作成日時を考慮した自動並び替え", 21, bullet=True)

# ============ Slide 4: デモ動画 ============
s = prs.slides.add_slide(BLANK)
content_header(s, 4, "デモ（実演動画）")
note = textbox(s, Inches(0.9), Inches(1.65), Inches(11.5), Inches(0.5))
add_para(note, "▶ 再生：ログイン → タスク作成 → カンバンでドラッグ → 検索・絞り込み（約27秒）", 16, color=MUTED, first=True)
vid = os.path.join(HERE, "demo.mp4")
poster = os.path.join(HERE, "poster.png")
vw, vh = Inches(8.0), Inches(4.5)   # 16:9
vleft = int((SW - vw) / 2)
vtop = Inches(2.2)
rect(s, vleft - Emu(40000), int(vtop) - Emu(40000), vw + Emu(80000), vh + Emu(80000), BORDER)
s.shapes.add_movie(vid, vleft, vtop, vw, vh, poster_frame_image=poster, mime_type="video/mp4")

# ============ Slide 5: 工夫① アーキテクチャ ============
s = prs.slides.add_slide(BLANK)
content_header(s, 5, "工夫① ストレージを差し替えられる設計")
tf = textbox(s, Inches(0.9), Inches(1.8), Inches(6.4), Inches(5.0))
add_para(tf, "リポジトリパターンで「保存方法」を抽象化。", 21, bold=True, first=True, space_after=14)
add_para(tf, "アプリ本体は具体的な DB ではなく共通インターフェースに依存", 19, bullet=True)
add_para(tf, "ローカル＝JSON ファイル、テスト＝メモリ、本番＝PostgreSQL", 19, bullet=True)
add_para(tf, "同じテストが全実装に通用（契約テスト）", 19, bullet=True)
add_para(tf, "環境変数で自動選択。設定ゼロで動く", 19, bullet=True)
# 右側に図解
bx = Inches(7.6); bw = Inches(5.0)
rect(s, bx, Inches(2.0), bw, Inches(0.8), INK)
t = textbox(s, bx, Inches(2.05), bw, Inches(0.7), MSO_ANCHOR.MIDDLE)
p = add_para(t, "UI / API（上位レイヤー）", 16, bold=True, color=WHITE, first=True); p.alignment = PP_ALIGN.CENTER
rect(s, bx + Inches(2.1), Inches(2.85), Inches(0.8), Inches(0.45), LIGHT)
t = textbox(s, bx, Inches(2.82), bw, Inches(0.5), MSO_ANCHOR.MIDDLE)
p = add_para(t, "↓ TaskRepository（共通I/F）", 14, color=MUTED, first=True); p.alignment = PP_ALIGN.CENTER
for i, (lbl, col) in enumerate([("JSON", LIGHT), ("メモリ", LIGHT), ("PostgreSQL", ACCENT)]):
    cw = Inches(1.55)
    cx = bx + i * (cw + Inches(0.18))
    rect(s, cx, Inches(3.5), cw, Inches(0.9), col, line=BORDER, line_w=Pt(1))
    t = textbox(s, cx, Inches(3.55), cw, Inches(0.8), MSO_ANCHOR.MIDDLE)
    p = add_para(t, lbl, 15, bold=True, color=(WHITE if col == ACCENT else INK), first=True)
    p.alignment = PP_ALIGN.CENTER

# ============ Slide 6: 工夫② 認証・マルチテナント ============
s = prs.slides.add_slide(BLANK)
content_header(s, 6, "工夫② 認証とマルチテナント（データ隔離）")
tf = body_box(s)
add_para(tf, "ユーザーごとにデータを完全に分離し、他人のタスクは一切見えない。", 22, bold=True, first=True, space_after=16)
add_para(tf, "Auth.js（NextAuth v5）でログイン。パスワードは scrypt でハッシュ化して保存", 20, bullet=True)
add_para(tf, "全データ操作を所有者IDでスコープ。取得・更新・削除も本人のものだけ", 20, bullet=True)
add_para(tf, "PostgreSQL では外部キー制約でも整合性を担保（二重の防御）", 20, bullet=True)
add_para(tf, "「他人のIDを指定しても 404」をテストで保証", 20, bullet=True)

# ============ Slide 7: 工夫③ 品質 ============
s = prs.slides.add_slide(BLANK)
content_header(s, 7, "工夫③ 実運用を意識した品質づくり")
tf = body_box(s)
add_para(tf, "「壊れていないこと」を仕組みで担保した。", 22, bold=True, first=True, space_after=16)
add_para(tf, "テスト 71 件（単体・契約・API結合・E2E）／カバレッジ 82% を CI で強制", 20, bullet=True)
add_para(tf, "GitHub Actions で 型チェック→Lint→テスト→ビルド を自動実行", 20, bullet=True)
add_para(tf, "セキュリティ：CSP 等のヘッダ、レート制限、CodeQL・Dependabot", 20, bullet=True)
add_para(tf, "観測性：構造化ログ・リクエストID・ヘルスチェック（/api/health）", 20, bullet=True)
add_para(tf, "Docker でどこでも起動できる（standalone ビルド）", 20, bullet=True)

# ============ Slide 8: 技術スタック ============
s = prs.slides.add_slide(BLANK)
content_header(s, 8, "技術スタック")
rows = [
    ("領域", "採用技術"),
    ("フレームワーク", "Next.js 14（App Router / Route Handlers）"),
    ("言語", "TypeScript（strict）"),
    ("認証", "Auth.js（NextAuth v5）＋ scrypt"),
    ("データベース", "PostgreSQL ＋ Drizzle ORM（本番）／ JSON・メモリ"),
    ("テスト", "Vitest ＋ PGlite ＋ Playwright（E2E）"),
    ("スタイル", "Tailwind CSS"),
    ("CI / 配布", "GitHub Actions・CodeQL・Docker"),
]
tbl_w = Inches(11.5)
tbl = s.shapes.add_table(len(rows), 2, Inches(0.9), Inches(1.85), tbl_w, Inches(4.6)).table
tbl.columns[0].width = Inches(3.4)
tbl.columns[1].width = tbl_w - Inches(3.4)
for ri, (a, b) in enumerate(rows):
    for ci, val in enumerate((a, b)):
        cell = tbl.cell(ri, ci)
        cell.margin_left = Inches(0.15); cell.margin_top = Inches(0.05); cell.margin_bottom = Inches(0.05)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.fill.solid()
        if ri == 0:
            cell.fill.fore_color.rgb = INK
        else:
            cell.fill.fore_color.rgb = WHITE if ri % 2 else LIGHT
        tf = cell.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]
        r = p.add_run(); r.text = val
        run_fmt(r, 16, bold=(ri == 0 or ci == 0),
                color=(WHITE if ri == 0 else INK))

# ============ Slide 9: 学び・今後 ============
s = prs.slides.add_slide(BLANK)
content_header(s, 9, "苦労した点・学んだこと・今後")
tf = body_box(s)
add_para(tf, "苦労した点・学び", 20, bold=True, color=ACCENT, first=True, space_after=10)
add_para(tf, "ドラッグ＆ドロップの状態管理。直感的に動かすため再描画の最適化に苦労した", 19, bullet=True)
add_para(tf, "「保存方法を後から変えられる設計」の大切さ（最初は JSON、後から DB へ）", 19, bullet=True)
add_para(tf, "テストがあると安心して機能追加できると実感した", 19, bullet=True)
add_para(tf, "今後やりたいこと", 20, bold=True, color=ACCENT, space_after=10)
add_para(tf, "通知・繰り返しタスク、チームでの共有、スマホ対応の強化", 19, bullet=True)

# ============ Slide 10: 参考文献 ============
s = prs.slides.add_slide(BLANK)
content_header(s, 10, "参考にしたサイト・参考文献")
tf = body_box(s)
add_para(tf, "Next.js 公式ドキュメント ― https://nextjs.org/docs", 19, bullet=True, first=True)
add_para(tf, "Auth.js（NextAuth）公式 ― https://authjs.dev", 19, bullet=True)
add_para(tf, "Drizzle ORM 公式 ― https://orm.drizzle.team", 19, bullet=True)
add_para(tf, "TypeScript 公式ハンドブック ― https://www.typescriptlang.org/docs", 19, bullet=True)
add_para(tf, "Tailwind CSS 公式 ― https://tailwindcss.com/docs", 19, bullet=True)
add_para(tf, "MDN Web Docs ― https://developer.mozilla.org/ja/", 19, bullet=True)
add_para(tf, "※ 各サイトの最終アクセス日：【記入してください】", 14, color=MUTED, space_after=4)

out = os.path.join(HERE, "発表スライド.pptx")
prs.save(out)
print("PPTX 出力:", out, f"{os.path.getsize(out)/1e6:.2f} MB")
